import io
import os
from datetime import datetime, timezone
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select
from app.models.analytics import SkillGapResult, CourseRelevanceScore, CurriculumRecommendation
from app.models.course import Course
from app.core.config import settings


async def generate_excel_report(db: AsyncSession) -> bytes:
    from openpyxl import Workbook
    wb = Workbook()

    # Skill Gaps sheet
    ws1 = wb.active
    ws1.title = "Skill Gaps"
    ws1.append(["Course ID", "Skill", "Category", "Similarity", "Mention %", "Gap Flag"])
    gaps = await db.execute(select(SkillGapResult).limit(1000))
    for g in gaps.scalars():
        ws1.append([
            str(g.course_id), g.skill_name, g.skill_category or "",
            g.max_similarity_score or 0, g.alumni_mention_pct or 0, g.gap_flag,
        ])

    # Relevance Scores sheet
    ws2 = wb.create_sheet("Relevance Scores")
    ws2.append(["Course ID", "Relevance Score", "Alumni Avg", "Skill Match", "CO Attainment"])
    scores = await db.execute(select(CourseRelevanceScore).limit(500))
    for s in scores.scalars():
        ws2.append([
            str(s.course_id), s.relevance_score or 0,
            s.alumni_relevance_avg or 0, s.skill_match_score or 0,
            s.co_attainment_avg or "",
        ])

    # Recommendations sheet
    ws3 = wb.create_sheet("Recommendations")
    ws3.append(["Course ID", "Type", "Topic", "Evidence", "Priority", "Status"])
    recs = await db.execute(select(CurriculumRecommendation).limit(500))
    for r in recs.scalars():
        ws3.append([
            str(r.course_id), r.recommendation_type.value, r.target_topic,
            r.evidence_summary or "", r.priority_score, r.status.value,
        ])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


async def generate_pdf_report(db: AsyncSession, course_id=None) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.lib import colors
    from reportlab.lib.units import cm

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4)
    styles = getSampleStyleSheet()
    story = []

    story.append(Paragraph("EduFeedback AI — Analytics Report", styles["Title"]))
    story.append(Paragraph(f"Generated: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')}", styles["Normal"]))
    story.append(Spacer(1, 1 * cm))

    # Recommendations section
    story.append(Paragraph("Top Curriculum Recommendations", styles["Heading2"]))
    query = select(CurriculumRecommendation).order_by(CurriculumRecommendation.priority_score.desc()).limit(20)
    if course_id:
        query = query.where(CurriculumRecommendation.course_id == course_id)
    recs_result = await db.execute(query)
    recs = recs_result.scalars().all()

    if recs:
        table_data = [["Type", "Topic", "Priority", "Status"]]
        for r in recs:
            table_data.append([
                r.recommendation_type.value,
                r.target_topic[:50],
                f"{r.priority_score:.2f}",
                r.status.value,
            ])
        tbl = Table(table_data, colWidths=[3*cm, 9*cm, 3*cm, 3*cm])
        tbl.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1D4ED8")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#EFF6FF")]),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.lightgrey),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
        ]))
        story.append(tbl)
    else:
        story.append(Paragraph("No recommendations found. Run analytics first.", styles["Normal"]))

    doc.build(story)
    buf.seek(0)
    return buf.read()
